import io
import os
import re
import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional
import docx
import httpx


class PDFConversionError(Exception):
    pass


class GotenbergPDFGenerator:
    """
    Fills official UK qualification Assessment Marking Sheet DOCX templates
    and converts them to PDF feedback reports via Gotenberg.

    Key improvement: per-criterion feedback is now stored individually so that
    criterion 1.1, 1.2, 2.1, 2.2, 3.1, 3.2, 3.3 each get their own distinct
    assessor comment — no more duplicated feedback blocks.
    """

    def __init__(self, gotenberg_url: str = None):
        self.gotenberg_url = gotenberg_url or os.getenv("GOTENBERG_URL", "http://gotenberg:3000")
        self.templates_dir = Path("/app/templates")

    def fill_feedback_template(
        self,
        template_bytes: bytes,
        student_name: str,
        student_id: str,
        unit_code: str,
        verdict: str,
        tasks: List[Dict[str, Any]],
        overall_comment: str = "",
        format_warnings: List[str] = None,
        rubric: Dict[str, Any] = None,
        per_criterion: Dict[str, Any] = None,
        ai_detection: Dict[str, Any] = None,
        plagiarism: Dict[str, Any] = None,
        review_required: bool = False,
    ) -> bytes:
        """
        Fill the feedback DOCX template with grading results.

        per_criterion: flat dict keyed by criterion ID → {"verdict": ..., "feedback": ...}
        e.g. {"1.1": {"verdict": "pass", "feedback": "..."}, "1.2": {...}, ...}
        """
        doc = docx.Document(io.BytesIO(template_bytes))
        is_pass = str(verdict).lower() == "pass"
        format_warnings = format_warnings or []
        rubric = rubric or {}
        per_criterion = per_criterion or {}
        ai_detection = ai_detection or {}
        plagiarism = plagiarism or {}

        # Build a flat ordered list of rubric criteria: [{id, task, description}, ...]
        all_criteria: List[Dict[str, Any]] = []
        for lo in rubric.get("learning_outcomes", []):
            for c in lo.get("criteria", []):
                all_criteria.append({
                    "id":          c.get("id", ""),
                    "task":        c.get("task", ""),
                    "description": c.get("description", ""),
                })

        # Build task lookup: task_number → task result dict
        task_by_num: Dict[int, Dict[str, Any]] = {
            t.get("task_number", 0): t for t in tasks
        }

        # ── Table 0: Learner Name ──────────────────────────────────────────────
        if len(doc.tables) > 0:
            t0 = doc.tables[0]
            if len(t0.rows) > 1 and t0.rows[1].cells:
                name_str = student_name
                if student_id and student_id not in ("N/A", ""):
                    name_str += f" — {student_id}"
                self._set_cell_text(t0.rows[1].cells[0], name_str)

        # ── Table 1: Pass / Refer tick ────────────────────────────────────────
        if len(doc.tables) > 1:
            t1 = doc.tables[1]
            if len(t1.rows) >= 2:
                if is_pass:
                    if len(t1.rows[0].cells) > 2:
                        self._set_cell_text(t1.rows[0].cells[2], "✓")
                    if len(t1.rows[1].cells) > 2:
                        self._set_cell_text(t1.rows[1].cells[2], "")
                else:
                    if len(t1.rows[0].cells) > 2:
                        self._set_cell_text(t1.rows[0].cells[2], "")
                    if len(t1.rows[1].cells) > 2:
                        self._set_cell_text(t1.rows[1].cells[2], "✓")

        # ── Table 2: Learning Outcomes / Criteria Matrix ──────────────────────
        #
        # Strategy: for each row in the criteria table, look up the criterion ID
        # from the rubric list and find its specific per-criterion feedback.
        # This ensures every row gets UNIQUE, targeted assessor comments.
        #
        if len(doc.tables) > 2:
            t2 = doc.tables[2]
            crit_idx = 0  # sequential index into all_criteria

            for r_idx in range(1, len(t2.rows)):
                row = t2.rows[r_idx]
                if len(row.cells) < 4:
                    continue

                # ── Resolve which criterion this row belongs to ──
                crit_id = ""
                matched_verdict = "Refer" if not is_pass else "Pass"
                matched_feedback = ""
                matched_task: Optional[Dict[str, Any]] = None

                # Priority 1: Use rubric criteria list (most accurate)
                if all_criteria and crit_idx < len(all_criteria):
                    rubric_crit = all_criteria[crit_idx]
                    crit_id = rubric_crit.get("id", "")
                    crit_desc = rubric_crit.get("description", "")

                    # Find per-criterion feedback for this specific criterion ID
                    if crit_id and crit_id in per_criterion:
                        crit_data = per_criterion[crit_id]
                        crit_verdict_raw = crit_data.get("verdict", "refer")
                        matched_verdict = "Pass" if str(crit_verdict_raw).lower() == "pass" else "Refer"
                        matched_feedback = crit_data.get("feedback", "").strip()

                    # Also resolve the parent task for fallback
                    task_label = rubric_crit.get("task", "")
                    task_num_m = re.search(r'\d+', task_label)
                    if task_num_m:
                        matched_task = task_by_num.get(int(task_num_m.group(0)))

                    # If no per-criterion feedback found, fall back to task feedback
                    if not matched_feedback and matched_task:
                        # Try to get criterion-specific feedback from task's per_criterion
                        task_per_crit = matched_task.get("per_criterion", {})
                        if crit_id and crit_id in task_per_crit:
                            crit_data = task_per_crit[crit_id]
                            matched_verdict = "Pass" if str(crit_data.get("verdict", "refer")).lower() == "pass" else "Refer"
                            matched_feedback = crit_data.get("feedback", "").strip()

                    # Generate a criterion-specific fallback if still empty
                    if not matched_feedback:
                        matched_feedback = self._generate_criterion_fallback(
                            crit_id=crit_id,
                            crit_desc=crit_desc,
                            task=matched_task,
                            overall_pass=is_pass,
                        )

                else:
                    # Priority 2: Infer from cell text (no rubric available)
                    crit_col_text = row.cells[1].text.strip() if len(row.cells) > 1 else ""
                    cell_task_match = re.search(r'task\s*(\d+)', crit_col_text, re.I)
                    if cell_task_match:
                        matched_task = task_by_num.get(int(cell_task_match.group(1)))

                    if not matched_task and tasks:
                        matched_task = tasks[(r_idx - 1) % len(tasks)]

                    if matched_task:
                        crit_verdict_raw = matched_task.get("verdict", "refer")
                        matched_verdict = "Pass" if str(crit_verdict_raw).lower() == "pass" else "Refer"
                        matched_feedback = matched_task.get("feedback_text", "").strip()

                    if not matched_feedback:
                        matched_feedback = (
                            "The learner has addressed this criterion satisfactorily."
                            if is_pass else
                            "Further development is required for this criterion."
                        )

                self._set_cell_text(row.cells[2], matched_verdict)
                self._set_cell_text(row.cells[3], matched_feedback)
                crit_idx += 1

        # ── Table 3: Overall Comment ──────────────────────────────────────────
        if len(doc.tables) > 3:
            t3 = doc.tables[3]
            if t3.rows and t3.rows[0].cells:
                comment_text = overall_comment
                if review_required:
                    comment_text = (
                        "⚠ MAIN_ASSESSOR REVIEW REQUIRED — academic integrity warning. "
                        "The Pass/Refer decision has not been auto-failed.\n\n"
                        + comment_text
                    )
                if format_warnings:
                    comment_text += (
                        f"\n\nFormatting notes: {'; '.join(format_warnings)}."
                    )
                integrity_note = self._integrity_summary(ai_detection, plagiarism, review_required)
                if integrity_note:
                    comment_text += f"\n\n{integrity_note}"
                self._set_cell_text(
                    t3.rows[0].cells[0],
                    f"Overall Assessor Comments:\n\n{comment_text}"
                )

        # ── Table 4: Assessor Sign-off ────────────────────────────────────────
        if len(doc.tables) > 4:
            t4 = doc.tables[4]
            for row in t4.rows:
                cell_text = row.cells[0].text if row.cells else ""
                if "Assessor" in cell_text and len(row.cells) > 1:
                    self._set_cell_text(row.cells[1], "GradePro AI Verified Assessor")
                if "Signature" in cell_text and len(row.cells) > 3:
                    self._set_cell_text(row.cells[3], datetime.date.today().strftime("%d/%m/%Y"))

        self._append_integrity_section(doc, ai_detection, plagiarism, review_required)

        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()

    def _integrity_summary(self, ai_detection: dict, plagiarism: dict, review_required: bool) -> str:
        if not ai_detection and not plagiarism:
            return ""
        ai_score = float(ai_detection.get("score") or 0)
        ai_flagged = "FLAGGED" if ai_detection.get("flagged") else "clear"
        plag_pct = float(plagiarism.get("overall_similarity_pct") or 0)
        matches = plagiarism.get("top_matches") or []
        match_txt = "; ".join(
            f"{m.get('source', 'source')} ({float(m.get('similarity', 0)) * 100:.1f}%)"
            if float(m.get("similarity", 0)) <= 1
            else f"{m.get('source', 'source')} ({float(m.get('similarity', 0)):.1f}%)"
            for m in matches[:5]
        ) or "none"
        banner = " Review required by MAIN_ASSESSOR before verdict is final." if review_required else ""
        return (
            "Academic integrity (no plagiarism, no AI-generated content): "
            f"AI-detection score {ai_score:.2f} ({ai_flagged}); "
            f"plagiarism similarity {plag_pct:.1f}%; top matches: {match_txt}.{banner}"
        )

    def _append_integrity_section(self, doc, ai_detection: dict, plagiarism: dict, review_required: bool):
        if not ai_detection and not plagiarism:
            return
        doc.add_heading("Academic Integrity Checks", level=1)
        if review_required:
            banner = doc.add_paragraph()
            run = banner.add_run(
                "WARNING: MAIN_ASSESSOR review required. Integrity thresholds were crossed. "
                "The academic verdict was not auto-failed."
            )
            run.bold = True
        ai_score = float(ai_detection.get("score") or 0)
        doc.add_paragraph(
            f"AI-generated content score: {ai_score:.2f} "
            f"({'FLAGGED' if ai_detection.get('flagged') else 'not flagged'}) "
            f"[model {ai_detection.get('model_version', 'ai_detector_v2')}]"
        )
        plag_pct = float(plagiarism.get("overall_similarity_pct") or 0)
        doc.add_paragraph(
            f"Plagiarism overall similarity: {plag_pct:.1f}% "
            f"[model {plagiarism.get('model_version', 'plagiarism_v1')}]"
        )
        matches = plagiarism.get("top_matches") or []
        if matches:
            doc.add_paragraph("Top matched sources:")
            for match in matches[:5]:
                sim = float(match.get("similarity") or 0)
                pct = sim * 100 if sim <= 1 else sim
                doc.add_paragraph(f"- {match.get('source', 'unknown')}: {pct:.1f}%", style=None)

    def _set_cell_text(self, cell, text: str):
        """
        Safely set a table cell's text content, preserving the cell's paragraph
        formatting (font, size, alignment) where possible.
        """
        try:
            # Clear existing paragraphs except the first
            for para in cell.paragraphs[1:]:
                p = para._element
                p.getparent().remove(p)
            # Set text in first paragraph
            if cell.paragraphs:
                cell.paragraphs[0].text = text
            else:
                cell.text = text
        except Exception:
            try:
                cell.text = text
            except Exception:
                pass

    def _generate_criterion_fallback(
        self,
        crit_id: str,
        crit_desc: str,
        task: Optional[Dict[str, Any]],
        overall_pass: bool,
    ) -> str:
        """Generate a criterion-level fallback comment when per-criterion data is missing."""
        label = f"criterion {crit_id}: {crit_desc}" if crit_id else crit_desc

        if task and task.get("verdict") == "pass":
            task_fb = task.get("feedback_text", "")
            # Trim task feedback and make it criterion-specific
            snippet = task_fb[:150].rstrip() + "..." if len(task_fb) > 150 else task_fb
            return (
                f"The learner has satisfactorily addressed {label}. {snippet}"
                if snippet else
                f"The learner has satisfactorily addressed {label}."
            )
        elif task and task.get("verdict") == "refer":
            return (
                f"Further development is required for {label}. "
                f"Ensure this criterion is explicitly addressed with supporting evidence and theory."
            )
        elif overall_pass:
            return f"The learner has satisfactorily addressed {label}."
        else:
            return (
                f"Further development is required. Please ensure {label} is explicitly "
                f"addressed with supporting evidence and examples."
            )

    async def convert_docx_to_pdf(self, docx_bytes: bytes) -> bytes:
        async with httpx.AsyncClient(timeout=45.0) as client:
            try:
                resp = await client.post(
                    f"{self.gotenberg_url}/forms/libreoffice/convert",
                    files={
                        "files": (
                            "report.docx",
                            docx_bytes,
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    },
                )
                if resp.status_code == 200:
                    return resp.content
            except Exception:
                pass
        return b""
